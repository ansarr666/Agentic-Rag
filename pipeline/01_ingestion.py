"""01_INGESTION: Connectors for PDF, Word (.docx), HTML, Text (.txt, .md), and PostgreSQL."""

import os
import re
import logging
from pathlib import Path
from dataclasses import dataclass, field
from typing import List, Dict, Any, Optional

logger = logging.getLogger(__name__)

@dataclass
class LoadedDocument:
    """Represents a raw document loaded from local storage or database."""
    content: str
    source_path: str
    filename: str
    extension: str
    metadata: Dict[str, Any] = field(default_factory=dict)

class DocumentConnector:
    """Universal connector for files (PDF, DOCX, HTML, TXT, MD) and database sources."""

    def __init__(self, supported_extensions: Optional[List[str]] = None):
        self.supported_extensions = supported_extensions or [
            ".txt", ".pdf", ".md", ".docx", ".html", ".htm"
        ]

    def load_file(self, file_path: Path | str) -> Optional[LoadedDocument]:
        """Load a single document based on its extension."""
        path = Path(file_path).resolve()
        if not path.exists():
            logger.warning(f"File not found: {path}")
            return None

        ext = path.suffix.lower()
        if ext not in self.supported_extensions:
            return None

        try:
            if ext == ".pdf":
                return self._load_pdf(path)
            elif ext == ".docx":
                return self._load_docx(path)
            elif ext in [".html", ".htm"]:
                return self._load_html(path)
            else:
                return self._load_text(path)
        except Exception as e:
            logger.error(f"Unexpected error loading {path.name}: {e}")
            return None

    def _load_text(self, path: Path) -> LoadedDocument:
        """Load plain text or markdown file with multi-encoding fallback."""
        content = ""
        for enc in ["utf-8", "utf-8-sig", "latin-1", "cp1252"]:
            try:
                with open(path, "r", encoding=enc) as f:
                    content = f.read()
                break
            except (UnicodeDecodeError, UnicodeError):
                continue

        if not content:
            try:
                with open(path, "r", encoding="utf-8", errors="replace") as f:
                    content = f.read()
            except Exception as e:
                logger.warning(f"Could not read text file {path.name}: {e}")
                content = ""

        return LoadedDocument(
            content=content,
            source_path=str(path),
            filename=path.name,
            extension=path.suffix.lower(),
            metadata={"size_bytes": path.stat().st_size, "modified_at": path.stat().st_mtime}
        )

    def _load_pdf(self, path: Path) -> LoadedDocument:
        """Extract text from PDF pages with page-boundary metadata."""
        content = ""
        page_count = 0
        try:
            import pypdf
            reader = pypdf.PdfReader(str(path))
            page_count = len(reader.pages)
            pages = []
            for i, p in enumerate(reader.pages):
                page_text = p.extract_text() or ""
                pages.append(f"--- Page {i+1} ---\n{page_text}")
            content = "\n\n".join(pages)
            if not content.strip():
                logger.warning(f"PDF {path.name} contains no extractable text (may be scanned image/OCR required).")
        except ImportError:
            logger.warning("pypdf not installed. Install with: pip install pypdf")
            content = f"[PDF extraction requires pypdf: {path.name}]"
        except Exception as e:
            logger.warning(f"pypdf extraction failed for {path.name}: {e}")
            content = f"[PDF Text Extraction failed for {path.name}: {e}]"

        return LoadedDocument(
            content=content,
            source_path=str(path),
            filename=path.name,
            extension=".pdf",
            metadata={"size_bytes": path.stat().st_size, "page_count": page_count}
        )

    def _load_docx(self, path: Path) -> LoadedDocument:
        """Extract text from Microsoft Word (.docx) document including tables."""
        content = ""
        try:
            import docx
            doc = docx.Document(str(path))
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            table_texts = []
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells)
                    if row_text.strip():
                        table_texts.append(row_text)
            content = "\n\n".join(paragraphs + table_texts)
        except ImportError:
            logger.warning("python-docx not installed. Install with: pip install python-docx")
            content = f"[DOCX extraction requires python-docx: {path.name}]"
        except Exception as e:
            logger.warning(f"DOCX extraction failed for {path.name}: {e}")
            content = f"[DOCX Extraction failed for {path.name}: {e}]"

        return LoadedDocument(
            content=content,
            source_path=str(path),
            filename=path.name,
            extension=".docx",
            metadata={"size_bytes": path.stat().st_size}
        )

    def _load_html(self, path: Path) -> LoadedDocument:
        """Extract clean text from HTML file, stripping script and style tags."""
        content = ""
        try:
            from bs4 import BeautifulSoup
            with open(path, "r", encoding="utf-8", errors="ignore") as f:
                soup = BeautifulSoup(f.read(), "html.parser")
                for script in soup(["script", "style", "nav", "footer"]):
                    script.decompose()
                content = soup.get_text(separator="\n\n")
        except ImportError:
            try:
                with open(path, "r", encoding="utf-8", errors="ignore") as f:
                    raw = f.read()
                    # Strip basic script and style blocks
                    raw = re.sub(r"<(script|style)[^>]*>.*?</\1>", "", raw, flags=re.DOTALL | re.IGNORECASE)
                    content = re.sub(r"<[^>]+>", " ", raw)
            except Exception as e:
                logger.warning(f"HTML fallback extraction failed for {path.name}: {e}")
                content = ""
        except Exception as e:
            logger.warning(f"HTML extraction failed for {path.name}: {e}")
            content = ""

        return LoadedDocument(
            content=content,
            source_path=str(path),
            filename=path.name,
            extension=path.suffix.lower(),
            metadata={"size_bytes": path.stat().st_size}
        )

    def load_directory(self, dir_path: Path | str, recursive: bool = True) -> List[LoadedDocument]:
        """Scan directory and load all supported documents."""
        folder = Path(dir_path).resolve()
        if not folder.exists() or not folder.is_dir():
            return []

        pattern = "**/*" if recursive else "*"
        documents = []
        for p in folder.glob(pattern):
            if p.is_file() and p.suffix.lower() in self.supported_extensions:
                doc = self.load_file(p)
                if doc and doc.content.strip():
                    documents.append(doc)

        return documents

class PostgreSQLConnector:
    """Ingests document records or rows directly from a company PostgreSQL table."""

    def __init__(self, connection_url: Optional[str] = None):
        self.connection_url = connection_url or os.getenv("DATABASE_URL")

    @staticmethod
    def _validate_identifier(name: str) -> str:
        """Sanitize SQL identifiers to strictly prevent SQL injection attacks."""
        if not re.match(r"^[a-zA-Z_][a-zA-Z0-9_]*$", name):
            raise ValueError(f"Invalid SQL identifier name: {name}")
        return name

    def fetch_documents(
        self,
        table_name: str = "documents",
        text_column: str = "content",
        id_column: str = "id",
        title_column: str = "title"
    ) -> List[LoadedDocument]:
        """Fetch rows from PostgreSQL table and convert to LoadedDocument objects."""
        if not self.connection_url:
            logger.info("DATABASE_URL not configured. Skipping PostgreSQL ingestion.")
            return []

        # Validate all column and table identifiers
        safe_table = self._validate_identifier(table_name)
        safe_id = self._validate_identifier(id_column)
        safe_title = self._validate_identifier(title_column)
        safe_text = self._validate_identifier(text_column)

        docs = []
        try:
            import psycopg2
            import psycopg2.extras
            with psycopg2.connect(self.connection_url) as conn:
                with conn.cursor(cursor_factory=psycopg2.extras.DictCursor) as cur:
                    query = f"SELECT {safe_id}, {safe_title}, {safe_text} FROM {safe_table} LIMIT 1000;"
                    cur.execute(query)
                    rows = cur.fetchall()
                    for r in rows:
                        doc_id = str(r[safe_id])
                        title = str(r.get(safe_title, doc_id))
                        body = str(r[safe_text])
                        if body.strip():
                            docs.append(LoadedDocument(
                                content=body,
                                source_path=f"postgres://{safe_table}/{doc_id}",
                                filename=f"pg_{safe_table}_{title}.txt",
                                extension=".pg",
                                metadata={"table": safe_table, "id": doc_id, "title": title}
                            ))
            logger.info(f"Successfully fetched {len(docs)} documents from PostgreSQL table {safe_table}.")
        except ImportError:
            logger.warning("psycopg2 not installed. Install with: pip install psycopg2-binary")
        except Exception as e:
            logger.error(f"Failed to fetch from PostgreSQL: {e}")

        return docs
